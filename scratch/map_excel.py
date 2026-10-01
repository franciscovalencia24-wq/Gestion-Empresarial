import pandas as pd
from openpyxl import load_workbook
import sys

def map_excel(file_path):
    wb = load_workbook(file_path, data_only=True)
    sheet = wb['Datos - solicitar a cliente']
    
    mapping = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None and isinstance(cell.value, str):
                val = cell.value.strip()
                if val:
                    print(f"{cell.coordinate}: {val}")

if __name__ == '__main__':
    map_excel('STONEX/PN - Datos apertura de cuenta Stonex (1).xlsx')
