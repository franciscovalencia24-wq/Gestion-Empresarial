import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def generar_excel_inventario_mfo(client_name="Cliente"):
    """
    Genera un archivo Excel con formato corporativo (Altus AI / FV Asesorías)
    para el levantamiento de Inventario Patrimonial (Propiedades, Cuentas, Sociedades, Pasivos).
    """
    wb = openpyxl.Workbook()
    
    # ----------------------------------------------------
    # ESTILOS CORPORATIVOS ELEGANTES
    # ----------------------------------------------------
    title_fill = PatternFill(start_color="0A2342", end_color="0A2342", fill_type="solid") # Dark Navy
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Navy Blue
    example_fill = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid") # Light Gray
    input_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid") # White
    
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    example_font = Font(name="Calibri", size=11, italic=True, color="595959")
    normal_font = Font(name="Calibri", size=11, color="000000")
    
    thin_border = Border(
        left=Side(style='thin', color="BFBFBF"),
        right=Side(style='thin', color="BFBFBF"),
        top=Side(style='thin', color="BFBFBF"),
        bottom=Side(style='thin', color="BFBFBF")
    )
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # ----------------------------------------------------
    # HOJA 1: INVENTARIO PATRIMONIAL
    # ----------------------------------------------------
    ws = wb.active
    ws.title = "Inventario Patrimonial"
    
    # Configurar anchos de columna
    cols = [('A', 20), ('B', 30), ('C', 20), ('D', 20), ('E', 30)]
    for col, width in cols:
        ws.column_dimensions[col].width = width
        
    # Título Principal
    ws.merge_cells('A1:E2')
    ws['A1'] = f"INVENTARIO PATRIMONIAL CONSOLIDADO MFO - {client_name.upper()}"
    ws['A1'].fill = title_fill
    ws['A1'].font = title_font
    ws['A1'].alignment = align_center
    
    ws['A3'] = "Por favor, complete los espacios en blanco. Los ejemplos en cursiva gris son solo de referencia."
    ws.merge_cells('A3:E3')
    ws['A3'].font = Font(name="Calibri", size=11, italic=True, color="333333")
    
    row_idx = 5
    
    def add_section(title, headers, examples):
        nonlocal row_idx
        # Título Sección
        ws.merge_cells(f'A{row_idx}:E{row_idx}')
        ws[f'A{row_idx}'] = title
        ws[f'A{row_idx}'].fill = title_fill
        ws[f'A{row_idx}'].font = title_font
        ws[f'A{row_idx}'].alignment = align_left
        row_idx += 1
        
        # Cabeceras
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = thin_border
        row_idx += 1
        
        # Ejemplos
        for ex in examples:
            for col_idx, val in enumerate(ex, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = val
                cell.fill = example_fill
                cell.font = example_font
                cell.alignment = align_center
                cell.border = thin_border
            row_idx += 1
            
        # Filas vacías para input
        for _ in range(5):
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = input_fill
                cell.font = normal_font
                cell.border = thin_border
            row_idx += 1
            
        row_idx += 2
        
    add_section(
        "1. PROPIEDADES INMOBILIARIAS (CBR)",
        ["Rol / Comuna", "Dirección", "Destino (Hab/Com/Sitio)", "Avalúo Fiscal (Aprox)", "Deuda Hipotecaria Restante"],
        [["1234-56 / Las Condes", "Av. Apoquindo 1234, Dpto 501", "Habitacional", "$150.000.000", "$45.000.000 (Banco BCI)"]]
    )
    
    add_section(
        "2. SOCIEDADES Y EMPRESAS (VPP)",
        ["RUT Sociedad", "Razón Social", "% Participación Propia", "Valor Estimado de la Sociedad (Total)", "Rubro / Comentarios"],
        [["76.123.456-7", "Inversiones Familiares SpA", "50%", "$500.000.000", "Holding Inversiones"]]
    )
    
    add_section(
        "3. CUENTAS BANCARIAS E INVERSIONES LÍQUIDAS",
        ["Institución / Banco", "Tipo (Cta Cte, Depósito, FFMM)", "N° Cuenta", "Saldo Promedio / Monto", "Moneda (CLP, USD, EUR)"],
        [["Banco de Chile", "Fondo Mutuo Accionario", "123456789", "$85.000.000", "CLP"], ["Pershing LLC", "Cuenta de Inversión", "XXX-123", "USD 120.000", "USD"]]
    )
    
    add_section(
        "4. VEHÍCULOS Y OTROS ACTIVOS DE VALOR",
        ["Tipo de Activo", "Marca / Modelo / Patente", "Año", "Valor Comercial Estimado", "Asegurado (Sí/No)"],
        [["Vehículo SUV", "Volvo XC90 - AB CD 12", "2023", "$45.000.000", "Sí"]]
    )
    
    add_section(
        "5. PASIVOS Y DEUDAS IMPORTANTES",
        ["Institución Acreedora", "Tipo de Deuda (Consumo, Comercial)", "Monto Total Adeudado", "Tasa / Plazo Restante", "Comentarios"],
        [["Banco Santander", "Crédito de Consumo", "$12.000.000", "1.2% / 24 meses", "Crédito automotriz"]]
    )
    
    # Return as bytes
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
